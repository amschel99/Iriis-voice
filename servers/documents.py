"""documents MCP server: create and read office documents inside the allowlisted folders.

Create tools: create_document (.docx), create_spreadsheet (.xlsx), create_pdf, create_presentation (.pptx).
Edit tool: append_to_spreadsheet. Read tool: read_document (.docx, .xlsx, .csv, .pptx, .pdf, .txt, .md).
Documents are written from simple Markdown: # headings, paragraphs, - bullets, 1. numbered lists, **bold**,
*italic*, | pipe | tables | and images on their own line: ![caption](picture.png). Existing files are never replaced unless overwrite=true.
Libraries: python-docx, openpyxl, python-pptx (MIT), reportlab, pypdf (BSD).
"""
from __future__ import annotations

import csv
import re
from pathlib import Path

from mcp.server.mcpserver.exceptions import ToolError

from servers._common import make_server
from servers.files import _resolve

mcp = make_server("documents")

MAX_READ_CHARS = 8000
MAX_SHEET_ROWS = 200
DOC_EXT = {".docx", ".doc", ".odt", ".rtf", ".txt", ".md", ".pdf", ".xlsx", ".xlsm", ".xls", ".csv", ".ods",
           ".pptx", ".ppt", ".odp"}
NUMBER = re.compile(r"^-?(0|[1-9]\d*)(\.\d+)?$")
INLINE = re.compile(r"\*\*(.+?)\*\*|__(.+?)__|\*(.+?)\*|(?<!\w)_(.+?)_(?!\w)")


# ---------------------------------------------------------------- helpers

def resolve_image(ref: str) -> Path:
    """Where an ![caption](ref) image lives. Here: a file in the allowed folders. The cloud studio server swaps
    this for one that fetches its own stored charts and diagrams ("file:12")."""
    p = _resolve(ref)
    if not p.is_file():
        raise ToolError(f"Image not found: {ref}")
    return p


def _image(ref: str) -> Path:
    import servers.documents as me     # late lookup so a server that imports this module can replace it
    return me.resolve_image(ref)


def _target(filename: str, ext: str, overwrite: bool) -> Path:
    """Resolve a new file's path inside the allowed folders, with the right extension."""
    name = filename.strip()
    if not name:
        raise ToolError("A file name is needed.")
    suffix = Path(name).suffix.lower()
    if suffix != ext:
        name = str(Path(name).with_suffix(ext)) if suffix in DOC_EXT else name + ext
    p = _resolve(name)
    if p.exists() and not overwrite:
        raise ToolError(f"{p.name} already exists in {p.parent.name}; pass overwrite=true to replace it.\n{p}")
    p.parent.mkdir(parents=True, exist_ok=True)
    return p


def _saved(p: Path, detail: str) -> str:
    return f"Saved {p.name} in {p.parent.name} ({detail}).\n{p}"


def _plural(n: int, word: str) -> str:
    return f"{n} {word}{'' if n == 1 else 's'}"


def _runs(text: str) -> list[tuple[str, bool, bool]]:
    """Split inline Markdown into (text, bold, italic) runs."""
    out, pos = [], 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], False, False))
        bold = m.group(1) or m.group(2)
        out.append((bold, True, False) if bold else (m.group(3) or m.group(4), False, True))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], False, False))
    return out or [("", False, False)]


def _plain(text: str) -> str:
    return "".join(t for t, _, _ in _runs(text))


def _table_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def parse_markdown(md: str) -> list[tuple]:
    """Blocks: ("h", level, text), ("p", text), ("ul", [items]), ("ol", [items]), ("table", [rows]),
    ("img", ref, caption)."""
    blocks: list[tuple] = []
    para: list[str] = []

    def flush():
        if para:
            blocks.append(("p", " ".join(para)))
            para.clear()

    for raw in md.replace("\r\n", "\n").split("\n"):
        line = raw.strip()
        if not line:
            flush()
            continue
        if m := re.match(r"^!\[(.*?)\]\((.+?)\)$", line):
            flush()
            blocks.append(("img", m.group(2).strip(), m.group(1).strip()))
        elif m := re.match(r"^(#{1,6})\s+(.*)$", line):
            flush()
            blocks.append(("h", min(len(m.group(1)), 3), m.group(2).strip().rstrip("#").strip()))
        elif m := re.match(r"^[-*+]\s+(.*)$", line):
            flush()
            if blocks and blocks[-1][0] == "ul":
                blocks[-1][1].append(m.group(1))
            else:
                blocks.append(("ul", [m.group(1)]))
        elif m := re.match(r"^\d+[.)]\s+(.*)$", line):
            flush()
            if blocks and blocks[-1][0] == "ol":
                blocks[-1][1].append(m.group(1))
            else:
                blocks.append(("ol", [m.group(1)]))
        elif line.startswith("|"):
            flush()
            if re.match(r"^\|?[\s:|-]+\|?$", line) and "-" in line:
                continue  # the |---|---| separator row
            if blocks and blocks[-1][0] == "table":
                blocks[-1][1].append(_table_row(line))
            else:
                blocks.append(("table", [_table_row(line)]))
        else:
            para.append(line)
    flush()
    return blocks


def _cell(v):
    """Numbers stay numbers; numeric strings become numbers (but not ones with leading zeros, like phone numbers)."""
    if isinstance(v, str):
        s = v.strip()
        if NUMBER.match(s):
            return float(s) if "." in s else int(s)
    return v


def _page(text: str, start: int, max_chars: int) -> str:
    limit = max(500, min(max_chars, 20000))
    start = max(0, start)
    chunk = text[start:start + limit]
    more = (f"\n[... {len(text) - start - limit} more characters: read_document again with start={start + limit}]"
            if len(text) > start + limit else "")
    return chunk + more


# ---------------------------------------------------------------- create tools

@mcp.tool()
def create_document(filename: str, content: str, title: str = "", overwrite: bool = False) -> str:
    """Create a Word document (.docx) from simple Markdown: # headings, paragraphs, - bullets, 1. numbered
    lists, **bold**, *italic*, | pipe | tables |. A bare filename is saved in Documents."""
    from docx import Document

    p = _target(filename, ".docx", overwrite)
    doc = Document()
    if title:
        doc.core_properties.title = title
        doc.add_heading(title, 0)

    def fill(par, text):
        for t, b, i in _runs(text):
            r = par.add_run(t)
            r.bold, r.italic = b or None, i or None

    blocks = parse_markdown(content)
    for b in blocks:
        if b[0] == "h":
            doc.add_heading(_plain(b[2]), b[1])
        elif b[0] == "p":
            fill(doc.add_paragraph(), b[1])
        elif b[0] in ("ul", "ol"):
            for item in b[1]:
                fill(doc.add_paragraph(style="List Bullet" if b[0] == "ul" else "List Number"), item)
        elif b[0] == "img":
            from docx.shared import Inches
            doc.add_picture(str(_image(b[1])), width=Inches(6.0))
            if b[2]:
                cap = doc.add_paragraph()
                cap.add_run(_plain(b[2])).italic = True
        elif b[0] == "table":
            rows = b[1]
            cols = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=cols)
            t.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci in range(cols):
                    par = t.cell(ri, ci).paragraphs[0]
                    fill(par, f"**{row[ci]}**" if ri == 0 and ci < len(row) and row[ci] else (row[ci] if ci < len(row) else ""))
    doc.save(p)
    paras = sum(1 for b in blocks if b[0] not in ("table", "img"))
    tables = sum(1 for b in blocks if b[0] == "table")
    return _saved(p, _plural(paras, "section") + (f", {_plural(tables, 'table')}" if tables else ""))


@mcp.tool()
def create_spreadsheet(filename: str, sheets: dict[str, list[list]], overwrite: bool = False) -> str:
    """Create an Excel spreadsheet (.xlsx). sheets: {sheet name: list of rows}, each row a list of cells. The first
    row is the bold header. Numbers stay numbers; text starting with "=" is a formula (e.g. "=SUM(B2:B5)")."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    if not sheets:
        raise ToolError("sheets is empty: give at least one sheet, e.g. {\"Sheet1\": [[\"Item\", \"Cost\"]]}.")
    p = _target(filename, ".xlsx", overwrite)
    wb = Workbook()
    wb.remove(wb.active)
    total = 0
    for name, rows in sheets.items():
        ws = wb.create_sheet(str(name)[:31] or "Sheet")
        for row in rows or []:
            ws.append([_cell(v) for v in (row if isinstance(row, (list, tuple)) else [row])])
        total += len(rows or [])
        if rows:
            for c in ws[1]:
                c.font = Font(bold=True)
            ws.freeze_panes = "A2"
        _autosize(ws, get_column_letter)
    wb.save(p)
    return _saved(p, f"{_plural(len(sheets), 'sheet')}, {_plural(total, 'row')}")


def _autosize(ws, get_column_letter) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                n = 10 if isinstance(c.value, str) and c.value.startswith("=") else len(str(c.value))
                widths[c.column] = max(widths.get(c.column, 0), n)
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(w + 2, 8), 60)


@mcp.tool()
def create_pdf(filename: str, content: str, title: str = "", overwrite: bool = False) -> str:
    """Create a PDF from simple Markdown: # headings, paragraphs, - bullets, 1. numbered lists, **bold**,
    *italic*, | pipe | tables |. A bare filename is saved in Documents."""
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    p = _target(filename, ".pdf", overwrite)
    st = getSampleStyleSheet()

    def markup(text: str) -> str:
        out = []
        for t, b, i in _runs(text):
            t = escape(t)
            out.append(f"<b>{t}</b>" if b else f"<i>{t}</i>" if i else t)
        return "".join(out)

    story = []
    if title:
        story.append(Paragraph(escape(title), st["Title"]))
    blocks = parse_markdown(content)
    for b in blocks:
        if b[0] == "h":
            story.append(Paragraph(markup(b[2]), st[f"Heading{b[1]}"]))
        elif b[0] == "p":
            story.append(Paragraph(markup(b[1]), st["BodyText"]))
        elif b[0] in ("ul", "ol"):
            items = [ListItem(Paragraph(markup(x), st["BodyText"])) for x in b[1]]
            kind = {"bulletType": "bullet", "start": "•"} if b[0] == "ul" else {"bulletType": "1"}
            story.append(ListFlowable(items, leftIndent=14, **kind))
        elif b[0] == "img":
            from reportlab.platypus import Image
            img = Image(str(_image(b[1])))
            w = A4[0] - 4 * cm
            img.drawHeight, img.drawWidth = img.drawHeight * w / img.drawWidth, w
            if img.drawHeight > A4[1] - 6 * cm:
                h = A4[1] - 6 * cm
                img.drawWidth, img.drawHeight = img.drawWidth * h / img.drawHeight, h
            story.append(img)
            if b[2]:
                story.append(Paragraph(f"<i>{markup(b[2])}</i>", st["BodyText"]))
            story.append(Spacer(1, 8))
        elif b[0] == "table":
            rows = b[1]
            cols = max(len(r) for r in rows)
            data = [[Paragraph(markup(f"**{r[c]}**" if ri == 0 and c < len(r) and r[c] else (r[c] if c < len(r) else "")),
                               st["BodyText"]) for c in range(cols)] for ri, r in enumerate(rows)]
            t = Table(data, repeatRows=1, colWidths=[(A4[0] - 4 * cm) / cols] * cols)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8e8e8")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story += [t, Spacer(1, 8)]
    if not story:
        story.append(Paragraph("", st["BodyText"]))
    pdf = SimpleDocTemplate(str(p), pagesize=A4, title=title or p.stem, leftMargin=2 * cm, rightMargin=2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm)
    pdf.build(story)
    from pypdf import PdfReader
    return _saved(p, _plural(len(PdfReader(p).pages), "page"))


@mcp.tool()
def create_presentation(filename: str, slides: list[dict], overwrite: bool = False) -> str:
    """Create a PowerPoint presentation (.pptx). slides: list of {"title": "...", "bullets": ["...", ...],
    "image": "picture.png (optional)", "notes": "speaker notes (optional)"}."""
    from pptx import Presentation

    if not slides:
        raise ToolError("slides is empty: give at least one {\"title\": ..., \"bullets\": [...]}.")
    p = _target(filename, ".pptx", overwrite)
    prs = Presentation()
    for i, s in enumerate(slides):
        s = s if isinstance(s, dict) else {"title": str(s)}
        bullets = [str(b) for b in (s.get("bullets") or [])]
        # first slide with no bullets is a title slide; others use "Title and Content"
        layout = prs.slide_layouts[0 if i == 0 and not bullets else 1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = _plain(str(s.get("title", "")))
        body = slide.placeholders[1]
        if bullets:
            tf = body.text_frame
            tf.text = _plain(bullets[0])
            for b in bullets[1:]:
                tf.add_paragraph().text = _plain(b)
        elif layout == prs.slide_layouts[0] and s.get("subtitle"):
            body.text = str(s["subtitle"])
        else:
            body._element.getparent().remove(body._element)
        if s.get("image"):
            from pptx.util import Inches
            left = Inches(5.2) if bullets else Inches(1.0)
            slide.shapes.add_picture(str(_image(str(s["image"]))), left, Inches(1.6),
                                     width=Inches(4.5) if bullets else Inches(8.0))
            if bullets:
                body.width = Inches(4.6)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(s["notes"])
    prs.save(p)
    return _saved(p, _plural(len(slides), "slide"))


@mcp.tool()
def append_to_spreadsheet(path: str, sheet: str, rows: list[list]) -> str:
    """Add rows to the end of a sheet in an existing .xlsx file (the sheet is created if missing). Numbers stay
    numbers; text starting with "=" is a formula."""
    from openpyxl import load_workbook

    p = _resolve(path if Path(path).suffix else path + ".xlsx")
    if not p.is_file():
        raise ToolError(f"No spreadsheet at {p}; use create_spreadsheet first.")
    if p.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ToolError(f"{p.name} is not an .xlsx spreadsheet.")
    wb = load_workbook(p, keep_vba=p.suffix.lower() == ".xlsm")
    created = sheet not in wb.sheetnames
    ws = wb.create_sheet(sheet[:31]) if created else wb[sheet]
    for row in rows:
        ws.append([_cell(v) for v in (row if isinstance(row, (list, tuple)) else [row])])
    wb.save(p)
    return (f"Added {_plural(len(rows), 'row')} to {'new sheet ' if created else ''}{ws.title} in {p.name} "
            f"(now {ws.max_row} rows).\n{p}")


# ---------------------------------------------------------------- read tool

def _docx_text(p: Path) -> str:
    from docx import Document

    doc = Document(p)
    out = []
    body = doc.element.body
    for el in body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = "".join(n.text or "" for n in el.iter() if n.tag.endswith("}t"))
            if t.strip():
                out.append(t)
        elif tag == "tbl":
            for tr in el.iter():
                if tr.tag.endswith("}tr"):
                    cells = ["".join(n.text or "" for n in tc.iter() if n.tag.endswith("}t"))
                             for tc in tr if tc.tag.endswith("}tc")]
                    out.append("\t".join(cells))
    return "\n".join(out)


def _xlsx_text(p: Path) -> str:
    from openpyxl import load_workbook

    values = load_workbook(p, data_only=True, read_only=True)
    formulas = load_workbook(p, read_only=True)
    out = []
    for ws in values.worksheets:
        fws = formulas[ws.title]
        out.append(f"[sheet {ws.title}]")
        n = 0
        for vrow, frow in zip(ws.iter_rows(values_only=True), fws.iter_rows(values_only=True)):
            if n >= MAX_SHEET_ROWS:
                out.append(f"[... more rows in {ws.title}]")
                break
            cells = [v if v is not None else f for v, f in zip(vrow, frow)]
            out.append("\t".join("" if c is None else str(c) for c in cells).rstrip("\t"))
            n += 1
    values.close()
    formulas.close()
    return "\n".join(out)


def _csv_text(p: Path) -> str:
    with p.open(newline="", encoding="utf-8-sig", errors="replace") as fh:
        rows = []
        for i, row in enumerate(csv.reader(fh)):
            if i >= MAX_SHEET_ROWS:
                rows.append("[... more rows]")
                break
            rows.append("\t".join(row))
    return "\n".join(rows)


def _pptx_text(p: Path) -> str:
    from pptx import Presentation

    out = []
    for i, slide in enumerate(Presentation(p).slides, 1):
        title = slide.shapes.title.text if slide.shapes.title is not None else ""
        out.append(f"[slide {i}] {title}".rstrip())
        for sh in slide.shapes:
            if sh == slide.shapes.title:
                continue
            if sh.has_text_frame and sh.text_frame.text.strip():
                out.append(sh.text_frame.text)
            elif getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    out.append("\t".join(c.text for c in r.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append(f"Notes: {slide.notes_slide.notes_text_frame.text}")
    return "\n".join(out)


def _pdf_text(p: Path) -> str:
    from pypdf import PdfReader

    r = PdfReader(p)
    return "\n\n".join(f"[page {i}]\n{(pg.extract_text() or '').strip()}" for i, pg in enumerate(r.pages, 1))


READERS = {".docx": _docx_text, ".xlsx": _xlsx_text, ".xlsm": _xlsx_text, ".csv": _csv_text,
           ".pptx": _pptx_text, ".pdf": _pdf_text}


@mcp.tool()
def read_document(path: str, max_chars: int = MAX_READ_CHARS, start: int = 0) -> str:
    """Read the text of a document: Word (.docx), Excel (.xlsx, .csv), PowerPoint (.pptx), PDF, .txt or .md.
    start: character offset, to continue a long document where the last read stopped."""
    p = _resolve(path)
    if not p.is_file():
        return f"Not a file: {p}"
    ext = p.suffix.lower()
    if ext in READERS:
        text = READERS[ext](p)
    elif ext in (".txt", ".md", ".markdown", ".rst", ".log"):
        text = p.read_text(encoding="utf-8", errors="replace")
    else:
        return f"Can't read {p.suffix or 'that'} files; supported: .docx, .xlsx, .csv, .pptx, .pdf, .txt, .md."
    if not text.strip():
        return f"{p.name} has no readable text (it may be scanned images).\n{p}"
    return f"{p.name}\n{p}\n\n{_page(text, start, max_chars)}"



@mcp.tool()
def save_from_cloud(file_id: int, folder: str = "", open_it: bool = True) -> str:
    """Save a chart, diagram or document Iriis made in the cloud (studio file id) onto this computer, in Documents
    or the given allowed folder, and open it."""
    import os

    import httpx
    url, key = os.environ.get("IRIIS_SERVER_URL", "").rstrip("/"), os.environ.get("IRIIS_API_KEY", "")
    if not url:
        raise ToolError("The computer doesn't know the cloud address (IRIIS_SERVER_URL); run iriis setup.")
    url = url.replace("wss://", "https://").replace("ws://", "http://")
    r = httpx.get(f"{url}/v1/files/{int(file_id)}", headers={"Authorization": f"Bearer {key}"}, timeout=60)
    if r.status_code != 200:
        raise ToolError(f"Couldn't fetch file {file_id} from the cloud (HTTP {r.status_code}).")
    import re as _re
    m = _re.search(r'filename="([^"]+)"', r.headers.get("content-disposition", ""))
    name = Path(m.group(1)).name if m else f"iriis_file_{file_id}"
    p = _resolve(str(Path(folder) / name) if folder else name)
    if p.exists():
        p = p.with_name(f"{p.stem}_{file_id}{p.suffix}")
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(r.content)
    if open_it and hasattr(os, "startfile"):
        os.startfile(p)  # noqa: S606  (the user's own file, opened with its usual app)
    return _saved(p, f"{len(r.content) / 1024:.0f} KB" + (", opened" if open_it else ""))


if __name__ == "__main__":
    mcp.run()
