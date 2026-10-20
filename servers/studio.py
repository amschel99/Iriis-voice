"""studio MCP server: charts, diagrams and documents made in the cloud, kept in the database, sent anywhere.

    make_chart       bar / line / area / pie / donut / scatter / histogram charts (Matplotlib)
    make_diagram     flowcharts, org charts, mind maps, process and system maps from Graphviz DOT
    make_document    PDF / Word / Excel / PowerPoint from simple Markdown, with charts and diagrams inside
    send_file        to the user's Telegram (photo or document)
    list_files       what's been made recently

Every file is stored in the `files` table (Railway Postgres), so it isn't tied to any computer; the desktop's
documents server can save one to the PC (save_from_cloud) and open it. All libraries are open source:
Matplotlib (PSF-style), Graphviz (EPL), python-docx / openpyxl / python-pptx (MIT), ReportLab (BSD).
"""
from __future__ import annotations

import base64
import json
import mimetypes
import os
import re
import tempfile
import time
from pathlib import Path

import httpx
from typing import Any

from servers._common import make_server
from iriis.storage import get_db

mcp = make_server("studio")
WORK = Path(tempfile.mkdtemp(prefix="iriis-studio-"))
os.environ["IRIIS_ALLOWED_DIRS"] = json.dumps([str(WORK)])    # the documents builders write here, then we store
MAX_BYTES = 15 * 1024 * 1024
for gv in (r"C:\Program Files\Graphviz\bin", r"C:\Program Files (x86)\Graphviz\bin"):
    if os.path.isdir(gv) and gv not in os.environ.get("PATH", ""):
        os.environ["PATH"] = os.environ.get("PATH", "") + os.pathsep + gv

PALETTE = ["#2E86C1", "#E67E22", "#27AE60", "#8E44AD", "#C0392B", "#16A085", "#F1C40F", "#7F8C8D", "#D35400", "#2C3E50"]


# ---------------------------------------------------------------------------- storage
def _slug(name: str, ext: str) -> str:
    base = re.sub(r"[^\w\- ]+", "", Path(name or "").stem).strip().replace(" ", "_")[:60] or f"iriis_{int(time.time())}"
    return f"{base}{ext}"


def store(name: str, data: bytes, mime: str | None = None) -> int:
    if len(data) > MAX_BYTES:
        raise ValueError(f"That file is {len(data) / 1e6:.1f} MB; the limit is 15 MB.")
    mime = mime or mimetypes.guess_type(name)[0] or "application/octet-stream"
    return get_db().insert("INSERT INTO files (name, mime, size, data, created) VALUES (?, ?, ?, ?, ?)",
                           (name, mime, len(data), base64.b64encode(data).decode(), time.time()))


def load(file_id: int) -> tuple[str, str, bytes]:
    rows = get_db().execute("SELECT name, mime, data FROM files WHERE id = ?", (int(file_id),))
    if not rows:
        raise ValueError(f"There's no file {file_id}. list_files shows what I've made.")
    name, mime, data = rows[0]
    return name, mime, base64.b64decode(data)


def _ready(fid: int, name: str, size: int, what: str) -> str:
    return (f"{what} ready: file {fid}, {name} ({size / 1024:.0f} KB). Send it with send_file, put it in a document "
            f"with ![caption](file:{fid}), or save it to the computer with save_from_cloud.")


# ---------------------------------------------------------------------------- charts
def render_chart(kind: str, title: str, labels: list, series: dict[str, list[float]], x_label: str = "",
                 y_label: str = "", stacked: bool = False, horizontal: bool = False, show_values: bool = False,
                 unit: str = "") -> bytes:
    import io

    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np

    kind = kind.lower().strip().replace(" ", "_")
    names = list(series)
    if not names:
        raise ValueError("series is empty: give e.g. {\"Sales\": [120, 150, 90]}.")
    vals = {k: [float(v) for v in series[k]] for k in names}
    n = max(len(v) for v in vals.values())
    labels = [str(x) for x in (labels or list(range(1, n + 1)))]
    plt.rcParams.update({"font.size": 11, "axes.spines.top": False, "axes.spines.right": False,
                         "axes.grid": True, "grid.alpha": 0.25, "axes.titleweight": "bold", "axes.titlesize": 14})
    fig, ax = plt.subplots(figsize=(9, 5.2), dpi=160)
    def fmt(v: float) -> str:
        return f"{unit}{v:,.0f}" if abs(v) >= 100 or float(v).is_integer() else f"{unit}{v:,.1f}"
    if kind in ("pie", "donut"):
        k = names[0]
        wedges, *_ = ax.pie(vals[k], labels=labels, autopct="%1.0f%%", startangle=90, colors=PALETTE[:len(labels)],
                            wedgeprops={"width": 0.45} if kind == "donut" else None, pctdistance=0.78 if kind == "donut" else 0.6)
        ax.axis("equal")
        ax.grid(False)
    elif kind in ("bar", "column", "stacked_bar", "barh"):
        stacked = stacked or kind == "stacked_bar"
        horizontal = horizontal or kind == "barh"
        x = np.arange(len(labels))
        width = 0.8 if stacked or len(names) == 1 else 0.8 / len(names)
        bottom = np.zeros(len(labels))
        for i, k in enumerate(names):
            v = np.array(vals[k] + [0] * (len(labels) - len(vals[k])))
            pos = x if stacked or len(names) == 1 else x - 0.4 + width * (i + 0.5)
            bars = (ax.barh(pos, v, width, left=bottom, label=k, color=PALETTE[i % 10]) if horizontal
                    else ax.bar(pos, v, width, bottom=bottom, label=k, color=PALETTE[i % 10]))
            if show_values:
                ax.bar_label(bars, labels=[fmt(t) for t in v], fontsize=9, padding=2,
                             label_type="center" if stacked else "edge")
            if stacked:
                bottom = bottom + v
        (ax.set_yticks if horizontal else ax.set_xticks)(x, labels, **({} if horizontal else {"rotation": 30 if len(labels) > 6 else 0}))
        ax.grid(axis="x" if not horizontal else "y", visible=False)
    elif kind in ("line", "area"):
        x = np.arange(len(labels))
        bottom = np.zeros(len(labels))
        for i, k in enumerate(names):
            v = np.array(vals[k] + [np.nan] * (len(labels) - len(vals[k])))
            if kind == "area":
                top = bottom + np.nan_to_num(v) if stacked else np.nan_to_num(v)
                ax.fill_between(x, bottom if stacked else 0, top, alpha=0.35, color=PALETTE[i % 10], label=k)
                ax.plot(x, top, color=PALETTE[i % 10], lw=1.8)
                bottom = top if stacked else bottom
            else:
                ax.plot(x, v, marker="o" if len(labels) <= 24 else None, lw=2.2, color=PALETTE[i % 10], label=k)
                if show_values:
                    for xi, yi in zip(x, v):
                        if not np.isnan(yi):
                            ax.annotate(fmt(yi), (xi, yi), textcoords="offset points", xytext=(0, 6), ha="center", fontsize=8)
        step = max(1, len(labels) // 12)
        ax.set_xticks(x[::step], labels[::step], rotation=30 if len(labels) > 8 else 0)
    elif kind == "scatter":
        for i, k in enumerate(names):
            pts = series[k]
            if pts and isinstance(pts[0], (list, tuple)):
                xs, ys = zip(*pts)
            else:
                xs, ys = [float(l) for l in labels[:len(pts)]], pts
            ax.scatter(xs, ys, s=40, alpha=0.8, color=PALETTE[i % 10], label=k)
    elif kind in ("histogram", "hist"):
        for i, k in enumerate(names):
            ax.hist(vals[k], bins="auto", alpha=0.6 if len(names) > 1 else 0.9, color=PALETTE[i % 10], label=k)
    else:
        raise ValueError(f"Unknown chart kind '{kind}'. Use bar, stacked_bar, barh, line, area, pie, donut, scatter or histogram.")
    if title:
        ax.set_title(title)
    if x_label and kind not in ("pie", "donut"):
        ax.set_xlabel(x_label)
    if y_label and kind not in ("pie", "donut"):
        ax.set_ylabel(y_label)
    if kind not in ("pie", "donut") and unit:
        from matplotlib.ticker import FuncFormatter
        (ax.xaxis if horizontal else ax.yaxis).set_major_formatter(FuncFormatter(lambda v, _: f"{unit}{v:,.0f}"))
    if len(names) > 1 and kind not in ("pie", "donut"):
        ax.legend(frameon=False)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    return buf.getvalue()


@mcp.tool()
def make_chart(kind: str, title: str, series: dict[str, list[float]], labels: list[str] | None = None,
               x_label: str = "", y_label: str = "", stacked: bool = False, horizontal: bool = False,
               show_values: bool = False, unit: str = "", filename: str = "") -> str:
    """Draw a chart as a PNG. kind: bar, stacked_bar, barh, line, area, pie, donut, scatter, histogram.
    series: {name: values}, e.g. {"Sales": [120, 150, 90], "Costs": [80, 95, 70]}; labels: the categories or
    x values, e.g. ["Jan", "Feb", "Mar"] (pie/donut use the first series). unit: prefix for numbers, e.g. "KES ".
    Only plot numbers you actually have (from tools or the user); never invent data."""
    png = render_chart(kind, title, labels or [], series, x_label, y_label, stacked, horizontal, show_values, unit)
    name = _slug(filename or title, ".png")
    return _ready(store(name, png, "image/png"), name, len(png), "Chart")


# ---------------------------------------------------------------------------- diagrams
@mcp.tool()
def make_diagram(dot: str, filename: str = "", engine: str = "dot", fmt: str = "png") -> str:
    """Draw a diagram from Graphviz DOT, e.g. 'digraph { rankdir=LR; node [shape=box, style="rounded,filled",
    fillcolor="#EAF2F8"]; Order -> Pack -> Deliver; }'. Good for flowcharts, processes, org charts, systems,
    timelines and mind maps. engine: dot (flows, hierarchies), neato or fdp (networks), twopi or circo (mind maps,
    rings). fmt: png, svg or pdf."""
    import graphviz
    fmt = fmt.lower() if fmt.lower() in ("png", "svg", "pdf") else "png"
    src = dot.strip()
    if not re.match(r"^(strict\s+)?(di)?graph\b", src, re.I):
        src = "digraph {\n" + src + "\n}"
    if fmt == "png" and "dpi" not in src:
        src = re.sub(r"\{", "{\n  graph [dpi=160, fontname=\"Helvetica\"]; node [fontname=\"Helvetica\"]; "
                            "edge [fontname=\"Helvetica\"];", src, count=1)
    try:
        data = graphviz.Source(src, engine=engine if engine in ("dot", "neato", "fdp", "twopi", "circo", "sfdp") else "dot").pipe(format=fmt)
    except graphviz.backend.execute.CalledProcessError as e:
        err = (e.stderr or b"").decode(errors="replace").strip().splitlines()
        raise ValueError(f"The diagram has a mistake: {err[-1] if err else e}. Fix the DOT and try again.") from e
    except graphviz.backend.execute.ExecutableNotFound as e:
        raise ValueError("Graphviz isn't installed where I run, so I can't draw diagrams right now.") from e
    name = _slug(filename or "diagram", f".{fmt}")
    return _ready(store(name, data), name, len(data), "Diagram")


# ---------------------------------------------------------------------------- documents
def _cloud_image(ref: str) -> Path:
    """![caption](file:12) -> that stored chart/diagram, written to the work folder for the builders."""
    m = re.fullmatch(r"(?:file:)?(\d+)", ref.strip())
    if not m:
        raise ValueError(f"Images in cloud documents are my own files: write ![caption](file:12), not {ref}.")
    name, _, data = load(int(m.group(1)))
    p = WORK / f"img_{m.group(1)}_{name}"
    p.write_bytes(data)
    return p


@mcp.tool()
def make_document(kind: str, filename: str, content: str = "", title: str = "",
                  sheets: dict[str, list[list]] | None = None, slides: list[dict] | None = None) -> str:
    """Make a document in the cloud (works with the computer off). kind: pdf, docx (Word), xlsx (Excel) or pptx
    (PowerPoint). pdf/docx: content in simple Markdown (# headings, - bullets, **bold**, | tables |) and charts or
    diagrams on their own line as ![caption](file:12). xlsx: sheets {name: rows}. pptx: slides [{"title",
    "bullets", "image": "file:12", "notes"}]. Then send_file or save_from_cloud."""
    import servers.documents as D
    D.resolve_image = _cloud_image
    kind = kind.lower().strip(". ").replace("word", "docx").replace("excel", "xlsx").replace("powerpoint", "pptx")
    name = _slug(filename or title, f".{kind}")
    for old in WORK.glob(name):
        old.unlink()
    try:
        if kind == "pdf":
            D.create_pdf(name, content, title)
        elif kind == "docx":
            D.create_document(name, content, title)
        elif kind == "xlsx":
            D.create_spreadsheet(name, sheets or {})
        elif kind == "pptx":
            D.create_presentation(name, slides or [])
        else:
            raise ValueError("kind must be pdf, docx, xlsx or pptx.")
    except Exception as e:  # noqa: BLE001
        raise ValueError(str(e).split("\n")[0]) from e
    data = (WORK / name).read_bytes()
    return _ready(store(name, data), name, len(data), {"pdf": "PDF", "docx": "Word document", "xlsx": "Spreadsheet",
                                                      "pptx": "Presentation"}[kind])


# ---------------------------------------------------------------------------- delivery
@mcp.tool()
def send_file(file_id: int, to: str = "telegram", caption: str = "") -> str:
    """Send a file I made (chart, diagram, document) to the user. to: "telegram" (pictures arrive as photos,
    everything else as a document). To put it on the computer, use the documents tool save_from_cloud instead."""
    name, mime, data = load(file_id)
    if to.lower().strip() != "telegram":
        raise ValueError("I can send files to Telegram; for the computer use save_from_cloud.")
    token = os.environ.get("TELEGRAM_BOT_TOKEN", "")
    chats = [c.strip() for c in os.environ.get("TELEGRAM_CHAT_ID", "").split(",") if c.strip()]
    if not token or not chats:
        raise ValueError("Telegram isn't set up (TELEGRAM_BOT_TOKEN and TELEGRAM_CHAT_ID on Railway).")
    photo = mime in ("image/png", "image/jpeg") and len(data) < 9_500_000
    method, field = ("sendPhoto", "photo") if photo else ("sendDocument", "document")
    for chat in chats:
        r = httpx.post(f"https://api.telegram.org/bot{token}/{method}", timeout=60,
                       data={"chat_id": chat, "caption": caption[:1000]}, files={field: (name, data, mime)})
        if r.status_code != 200 or not r.json().get("ok"):
            raise ValueError(f"Telegram refused it: {r.json().get('description', r.status_code)}")
    return f"Sent {name} to Telegram."


@mcp.tool()
def list_files(limit: int = 15) -> str:
    """Charts, diagrams and documents I've made recently, newest first (id, name, size, when)."""
    rows = get_db().execute("SELECT id, name, size, created FROM files ORDER BY id DESC LIMIT ?", (max(1, min(limit, 50)),))
    if not rows:
        return "I haven't made any files yet."
    import datetime as dt
    from zoneinfo import ZoneInfo
    tz = ZoneInfo(os.environ.get("IRIIS_TIMEZONE", "Africa/Nairobi"))
    return "\n".join(f"file {i}: {n} ({s / 1024:.0f} KB, {dt.datetime.fromtimestamp(c, tz).strftime('%d %b %H:%M')})"
                     for i, n, s, c in rows)


# ---------------------------------------------------------------------------- maths plots and hand-drawn sketches
_MATH = {"sin", "cos", "tan", "arcsin", "arccos", "arctan", "arctan2", "sinh", "cosh", "tanh", "exp", "log", "log10",
         "log2", "sqrt", "abs", "floor", "ceil", "sign", "minimum", "maximum", "where", "clip", "pi", "e"}


def safe_expr(expr: str, **vars_) -> "Any":
    """Evaluate a maths expression over numpy arrays: numbers, + - * / ** %, comparisons, the functions above and
    the given variables only. No attributes, no indexing, no names outside that list."""
    import ast

    import numpy as np
    src = expr.replace("^", "**").strip()
    tree = ast.parse(src, mode="eval")
    allowed = (ast.Expression, ast.BinOp, ast.UnaryOp, ast.Compare, ast.BoolOp, ast.Call, ast.Name, ast.Load,
               ast.Constant, ast.Add, ast.Sub, ast.Mult, ast.Div, ast.Pow, ast.Mod, ast.FloorDiv, ast.USub, ast.UAdd,
               ast.Lt, ast.LtE, ast.Gt, ast.GtE, ast.Eq, ast.NotEq, ast.And, ast.Or, ast.IfExp)
    for node in ast.walk(tree):
        if not isinstance(node, allowed):
            raise ValueError(f"'{expr}': only maths is allowed ({type(node).__name__} isn't)")
        if isinstance(node, ast.Name) and node.id not in _MATH and node.id not in vars_:
            raise ValueError(f"'{expr}': unknown name '{node.id}'")
        if isinstance(node, ast.Call) and not (isinstance(node.func, ast.Name) and node.func.id in _MATH):
            raise ValueError(f"'{expr}': unknown function")
    env = {n: getattr(np, n) for n in _MATH if hasattr(np, n)}
    env["abs"] = np.abs
    env.update(vars_)
    with np.errstate(all="ignore"):
        return eval(compile(tree, "<expr>", "eval"), {"__builtins__": {}}, env)  # noqa: S307  (AST-checked above)


def _fig(hand: bool, size=(8, 6)):
    import logging

    import matplotlib
    matplotlib.use("Agg")
    logging.getLogger("matplotlib.font_manager").setLevel(logging.ERROR)   # no xkcd font installed: that's fine
    import matplotlib.pyplot as plt
    ctx = plt.xkcd(scale=1.2, length=90, randomness=2) if hand else plt.rc_context({"font.size": 11})
    return plt, ctx


@mcp.tool()
def plot_math(functions: list[dict] | None = None, x_range: list[float] | None = None, parametric: list[dict] | None = None,
              points: list[dict] | None = None, title: str = "", x_label: str = "x", y_label: str = "y",
              equal_aspect: bool = False, hand_drawn: bool = False, y_range: list[float] | None = None,
              filename: str = "") -> str:
    """Plot 2D maths: functions y = f(x), parametric curves, and labelled points, as a PNG.
    functions: [{"expr": "sin(x) * exp(-x/5)", "label": "damped"}]; x_range: [xmin, xmax];
    parametric: [{"x": "cos(t)", "y": "sin(t)", "t": [0, 6.283], "label": "circle"}];
    points: [{"x": 1, "y": 2, "label": "A"}]; y_range: [ymin, ymax] (optional; asymptotes are clipped). Allowed: + - * / ^, sin cos tan exp log sqrt abs pi e, etc.
    hand_drawn=true draws it like a whiteboard sketch. Good for explaining maths, physics and engineering."""
    import io

    import numpy as np
    lo, hi = (x_range or [-10, 10])[:2]
    plt, ctx = _fig(hand_drawn)
    with ctx:
        fig, ax = plt.subplots(figsize=(8.5, 6), dpi=150)
        x = np.linspace(float(lo), float(hi), 1200)
        ys = []
        for f in functions or []:
            y = np.broadcast_to(np.asarray(safe_expr(str(f["expr"]), x=x), dtype=float), x.shape).copy()
            y[~np.isfinite(y)] = np.nan
            ys.append(y)
        ylim = None
        if y_range:
            ylim = (float(y_range[0]), float(y_range[1]))
        elif ys and np.isfinite(np.concatenate(ys)).any():
            a, b = np.nanpercentile(np.concatenate(ys), [3, 97])      # robust: asymptotes don't squash the view
            pad = (b - a) * 0.2 or 1.0
            lo_all, hi_all = np.nanmin(np.concatenate(ys)), np.nanmax(np.concatenate(ys))
            ylim = (max(lo_all, a - pad) - pad * 0.1, min(hi_all, b + pad) + pad * 0.1)
        for i, (f, y) in enumerate(zip(functions or [], ys)):
            if ylim:                                   # break the line where it leaves the view (asymptotes)
                span = ylim[1] - ylim[0]
                y[(y < ylim[0] - span) | (y > ylim[1] + span)] = np.nan
            ax.plot(x, y, lw=2.2, color=PALETTE[i % 10], label=f.get("label") or f"y = {f['expr']}")
        for j, p in enumerate(parametric or []):
            t0, t1 = (p.get("t") or [0, 2 * np.pi])[:2]
            t = np.linspace(float(t0), float(t1), 1500)
            xs = np.broadcast_to(np.asarray(safe_expr(str(p["x"]), t=t), float), t.shape)
            ys = np.broadcast_to(np.asarray(safe_expr(str(p["y"]), t=t), float), t.shape)
            ax.plot(xs, ys, lw=2.2, color=PALETTE[(len(functions or []) + j) % 10], label=p.get("label"))
        for p in points or []:
            ax.scatter([float(p["x"])], [float(p["y"])], s=45, color="#C0392B", zorder=5)
            if p.get("label"):
                ax.annotate(str(p["label"]), (float(p["x"]), float(p["y"])), textcoords="offset points", xytext=(6, 6))
        if not hand_drawn:
            ax.grid(True, alpha=0.3)
            ax.axhline(0, color="#999", lw=0.8)
            ax.axvline(0, color="#999", lw=0.8)
        if equal_aspect:
            ax.set_aspect("equal", adjustable="datalim")
        ax.set_xlabel(x_label)
        ax.set_ylabel(y_label)
        if title:
            ax.set_title(title)
        if ylim:
            ax.set_ylim(*ylim)
        if ax.get_legend_handles_labels()[0]:
            ax.legend(frameon=False)
        fig.tight_layout()
        buf = io.BytesIO()
        fig.savefig(buf, format="png")
        plt.close(fig)
    png = buf.getvalue()
    name = _slug(filename or title or "plot", ".png")
    return _ready(store(name, png, "image/png"), name, len(png), "Plot")


@mcp.tool()
def make_sketch(elements: list[dict], title: str = "", hand_drawn: bool = True, filename: str = "") -> str:
    """Draw a picture yourself, by hand, like on a whiteboard: explain a setup, a layout, what the camera saw, a
    device, a room plan or an idea. The canvas is 100 wide x 70 high, (0, 0) bottom-left. elements, each with
    optional "color", "fill", "label", "width":
      {"shape": "rect", "x": 10, "y": 10, "w": 30, "h": 20}      {"shape": "circle", "x": 50, "y": 35, "r": 8}
      {"shape": "ellipse", "x": 50, "y": 35, "w": 30, "h": 12}    {"shape": "line", "points": [[0,0],[20,10]]}
      {"shape": "arrow", "from": [10, 10], "to": [40, 30]}       {"shape": "polygon", "points": [[..], ..]}
      {"shape": "text", "x": 50, "y": 65, "text": "Fridge", "size": 14}
    Labels sit at the shape's centre (arrows: midpoint). Colours: names or hex. hand_drawn=false for clean lines."""
    import io

    from matplotlib import patches
    plt, ctx = _fig(hand_drawn)
    with ctx:
        fig, ax = plt.subplots(figsize=(10, 7), dpi=150)
        ax.set_xlim(0, 100)
        ax.set_ylim(0, 70)
        ax.set_aspect("equal")
        ax.axis("off")
        for i, el in enumerate(elements or []):
            kind = str(el.get("shape", "")).lower()
            color = el.get("color") or "#2C3E50"
            fill = el.get("fill")
            lw = float(el.get("width", 2))
            kw = {"edgecolor": color, "facecolor": fill or "none", "lw": lw, "alpha": 0.9}
            centre = None
            if kind in ("rect", "rectangle", "box"):
                x, y, w, h = (float(el[k]) for k in ("x", "y", "w", "h"))
                ax.add_patch(patches.FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.4,rounding_size=1.2", **kw))
                centre = (x + w / 2, y + h / 2)
            elif kind == "circle":
                x, y, r = float(el["x"]), float(el["y"]), float(el["r"])
                ax.add_patch(patches.Circle((x, y), r, **kw))
                centre = (x, y)
            elif kind == "ellipse":
                x, y = float(el["x"]), float(el["y"])
                ax.add_patch(patches.Ellipse((x, y), float(el["w"]), float(el["h"]), **kw))
                centre = (x, y)
            elif kind in ("line", "polyline"):
                pts = [(float(a), float(b)) for a, b in el["points"]]
                ax.plot(*zip(*pts), color=color, lw=lw)
                centre = pts[len(pts) // 2]
            elif kind == "arrow":
                (x0, y0), (x1, y1) = el["from"], el["to"]
                ax.annotate("", xy=(float(x1), float(y1)), xytext=(float(x0), float(y0)),
                            arrowprops={"arrowstyle": "-|>", "color": color, "lw": lw, "mutation_scale": 18})
                centre = ((float(x0) + float(x1)) / 2, (float(y0) + float(y1)) / 2 + 2)
            elif kind == "polygon":
                ax.add_patch(patches.Polygon([(float(a), float(b)) for a, b in el["points"]], closed=True, **kw))
                pts = el["points"]
                centre = (sum(float(p[0]) for p in pts) / len(pts), sum(float(p[1]) for p in pts) / len(pts))
            elif kind == "text":
                ax.text(float(el["x"]), float(el["y"]), str(el.get("text", "")), ha="center", va="center",
                        fontsize=float(el.get("size", 13)), color=color)
                continue
            else:
                raise ValueError(f"element {i}: unknown shape '{kind}' (rect, circle, ellipse, line, arrow, polygon, text)")
            if el.get("label") and centre:
                ax.text(centre[0], centre[1], str(el["label"]), ha="center", va="center", fontsize=12, color="#111")
        if title:
            ax.set_title(title, fontsize=16)
        fig.tight_layout()
        buf = io.BytesIO()
        fig.savefig(buf, format="png", facecolor="white")
        plt.close(fig)
    png = buf.getvalue()
    name = _slug(filename or title or "sketch", ".png")
    return _ready(store(name, png, "image/png"), name, len(png), "Sketch")


# ---------------------------------------------------------------------------- the camera and what the user sends
def _vision() -> tuple[str, str]:
    url, key = os.environ.get("VISION_MCP_URL", "").strip(), os.environ.get("VISION_ACCESS_KEY", "").strip()
    return (url.rstrip("/").removesuffix("/mcp"), key) if url and key else ("", "")


@mcp.tool()
def camera_photo(filename: str = "") -> str:
    """Take a photo from the live camera (the user's phone on a video call) and keep it as a file, so you can put
    it in a document (![caption](file:N)), send it, or look at it closely with inspect_file."""
    base, key = _vision()
    if not base:
        raise ValueError("The camera isn't connected (VISION_MCP_URL / VISION_ACCESS_KEY).")
    r = httpx.get(f"{base}/api/frame", headers={"Authorization": f"Bearer {key}"}, timeout=20)
    if r.status_code == 404:
        raise ValueError("No phone is streaming right now; ask the user to open the video call.")
    r.raise_for_status()
    name = _slug(filename or f"camera_{time.strftime('%Y%m%d_%H%M%S')}", ".jpg")
    return _ready(store(name, r.content, "image/jpeg"), name, len(r.content), "Photo")


_GEMINI_CACHE: dict[str, Any] = {}


def _gemini_models(base: str, key: str) -> list[str]:
    """GEMINI_VISION_MODEL (default gemini-flash-latest) first, then up to 3 current Flash models this key can use,
    found from the model list once an hour; models that answered 404 (retired) are left out."""
    first = os.environ.get("GEMINI_VISION_MODEL", "gemini-flash-latest")
    if time.time() - _GEMINI_CACHE.get("at", 0) > 3600:
        found = []
        try:
            r = httpx.get(f"{base}/models", headers={"Authorization": f"Bearer {key}"}, timeout=12)
            ids = [str(m.get("id", "")).removeprefix("models/") for m in r.json().get("data", [])] if r.status_code == 200 else []
            skip = ("tts", "image", "live", "embed", "audio", "native", "thinking", "exp", "learnlm", "gemma")
            flash = [m for m in ids if "flash" in m and not any(k in m for k in skip)]
            num = lambda m: float((re.findall(r"(\d+(?:\.\d+)?)", m) or ["0"])[0])  # noqa: E731
            found = sorted(flash, key=lambda m: (-num(m), "lite" in m, "preview" in m, len(m)))[:3]
        except Exception:  # noqa: BLE001
            pass
        _GEMINI_CACHE.update(at=time.time(), found=found)
    dead = _GEMINI_CACHE.setdefault("dead", set())
    return [m for m in dict.fromkeys([first, *_GEMINI_CACHE.get("found", [])]) if m not in dead]


def vision_describe(images: list[bytes], question: str, max_tokens: int = 700) -> str:
    """Ask a vision model about pictures: Gemini (free tier) first, then Groq. Both are OpenAI-compatible."""
    providers = []
    if os.environ.get("GEMINI_API_KEY"):
        gbase, gkey = "https://generativelanguage.googleapis.com/v1beta/openai", "".join(os.environ["GEMINI_API_KEY"].split())
        for model in _gemini_models(gbase, gkey):      # current Flash models (names retire; overloaded ones vary)
            providers.append((gbase, gkey, model, 8))
    if os.environ.get("GROQ_API_KEY"):
        providers.append(("https://api.groq.com/openai/v1", "".join(os.environ["GROQ_API_KEY"].split()),
                          os.environ.get("GROQ_VISION_MODEL", "qwen/qwen3.8-27b"), 3))
    if not providers:
        raise ValueError("No vision model is configured (GEMINI_API_KEY or GROQ_API_KEY).")
    errors = []
    for i, (base, key, model, limit) in enumerate(providers):
        last = i == len(providers) - 1
        parts = [{"type": "text", "text": question}] + [
            {"type": "image_url", "image_url": {"url": "data:image/jpeg;base64," + base64.b64encode(_jpeg(b)).decode()}}
            for b in images[:limit]]
        try:
            for attempt in range(3):
                r = httpx.post(f"{base}/chat/completions", timeout=90, headers={"Authorization": f"Bearer {key}"},
                               json={"model": model, "max_tokens": max_tokens, "temperature": 0.2,
                                     "messages": [{"role": "user", "content": parts}]})
                if r.status_code == 404 and "gemini" in model:
                    _GEMINI_CACHE.setdefault("dead", set()).add(model)     # retired: never again
                # busy or overloaded: move straight on to the next model; only the last one waits and retries
                if r.status_code not in (429, 503) or attempt == 2 or not last:
                    break
                try:                                    # busy: wait as asked (capped), then try again
                    wait = min(float(r.headers.get("retry-after", "") or 6 * (attempt + 1)), 20)
                except ValueError:
                    wait = 6 * (attempt + 1)
                time.sleep(wait)
            r.raise_for_status()
            text = r.json()["choices"][0]["message"]["content"] or ""
            return re.sub(r"<think>.*?</think>", "", text, flags=re.S).strip()
        except Exception as e:  # noqa: BLE001
            errors.append(f"{model}: {str(e)[:120]}")
    raise ValueError("The vision models didn't answer: " + "; ".join(errors))


def _jpeg(data: bytes, side: int = 1280) -> bytes:
    """Any image -> a JPEG no bigger than `side` px (smaller uploads, every provider accepts it)."""
    import io
    try:
        from PIL import Image        # comes with Matplotlib
        im = Image.open(io.BytesIO(data)).convert("RGB")
        im.thumbnail((side, side))
        out = io.BytesIO()
        im.save(out, "JPEG", quality=85)
        return out.getvalue()
    except Exception:  # noqa: BLE001  (not an image Pillow knows: send as is)
        return data


def contact_sheet(frames: list[tuple[float, bytes]], cols: int = 3, width: int = 1500) -> bytes:
    """Frames in a grid, left to right then top to bottom, each stamped with its time."""
    import io

    from PIL import Image, ImageDraw
    ims = [Image.open(io.BytesIO(b)).convert("RGB") for _, b in frames]
    cw = width // cols
    ch = int(cw * ims[0].height / max(1, ims[0].width))
    rows = (len(ims) + cols - 1) // cols
    sheet = Image.new("RGB", (cw * cols, ch * rows), "black")
    d = ImageDraw.Draw(sheet)
    for i, ((t, _), im) in enumerate(zip(frames, ims)):
        x, y = (i % cols) * cw, (i // cols) * ch
        sheet.paste(im.resize((cw - 4, ch - 4)), (x + 2, y + 2))
        label = f"{i + 1}  {int(t // 60)}:{int(t % 60):02d}"
        d.rectangle((x + 6, y + 6, x + 12 + 9 * len(label), y + 30), fill="black")
        d.text((x + 10, y + 10), label, fill="white")
    out = io.BytesIO()
    sheet.save(out, "JPEG", quality=82)
    return out.getvalue()


def _video_frames(data: bytes, n: int = 6) -> tuple[list[tuple[float, bytes]], float]:
    """n frames spread over the video, always including the first and the very last moment."""
    import io

    import av

    def jpeg(fr) -> bytes:
        out = io.BytesIO()
        fr.to_image().save(out, "JPEG", quality=80)
        return out.getvalue()
    frames, duration, last = [], 0.0, None
    with av.open(io.BytesIO(data)) as c:
        st = next((s for s in c.streams if s.type == "video"), None)
        if st is None:
            return [], 0.0
        duration = float(c.duration / 1e6) if c.duration else 0.0
        every = max(duration / max(1, n - 1), 0.5) if duration else 1.0
        nxt = 0.0
        for fr in c.decode(st):
            t = float(fr.time or 0)
            last = fr
            if t + 1e-6 >= nxt and len(frames) < n - 1:
                frames.append((t, jpeg(fr)))
                nxt = t + every
    if last is not None and (not frames or float(last.time or 0) > frames[-1][0] + 0.1):
        frames.append((float(last.time or 0), jpeg(last)))
    return frames, duration


def _transcribe(name: str, data: bytes, mime: str) -> str:
    """Her own speech-to-text (the server's /v1/stt), so audio and video soundtracks become words."""
    port, key = os.environ.get("PORT", "8000"), os.environ.get("IRIIS_API_KEY", "")
    r = httpx.post(f"http://127.0.0.1:{port}/v1/stt", files={"file": (name, data, mime)}, timeout=180,
                   headers={"Authorization": f"Bearer {key}"} if key else {})
    r.raise_for_status()
    return (r.json().get("text") or "").strip()


@mcp.tool()
def inspect_file(file_id: int, question: str = "") -> str:
    """Open a file the user sent in the chat (or one you made) and read it: PDFs, Word, Excel, PowerPoint, text,
    CSV and code are read as text; pictures are looked at by a vision model; audio is transcribed; videos are
    watched (key frames) and listened to. question: what to look for, e.g. "what's wrong in this wiring?"."""
    import io
    name, mime, data = load(file_id)
    low = name.lower()
    q = question.strip() or "Describe this in detail: what it shows, any text, numbers, problems or things worth noting."
    if mime.startswith("image/") or low.endswith((".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".heic")):
        return f"{name} (picture): " + vision_describe([data], q)
    if mime.startswith("audio/") or low.endswith((".mp3", ".wav", ".ogg", ".oga", ".m4a", ".webm", ".flac", ".aac", ".opus")):
        return f"{name} (audio) transcript: " + (_transcribe(name, data, mime) or "(no speech heard)")
    if mime.startswith("video/") or low.endswith((".mp4", ".mov", ".mkv", ".avi", ".3gp")):
        frames, dur = _video_frames(data, n=9 if len(data) > 2_000_000 else 6)
        seen = vision_describe([contact_sheet(frames)], f"This image is a grid of {len(frames)} frames taken in order from a "
                               f"{dur:.0f} s video (numbered, with their times; left to right, then top to bottom). "
                               f"Describe what happens over time. {q}") if frames else ""
        try:
            heard = _transcribe(name, data, mime)
        except Exception:  # noqa: BLE001
            heard = ""
        return f"{name} (video, {dur:.0f} s). What it shows: {seen or '(no frames)'}\nWhat's said: {heard or '(no speech)'}"
    text = ""
    if low.endswith(".pdf") or mime == "application/pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(f"[page {i + 1}] {(p.extract_text() or '').strip()}" for i, p in enumerate(reader.pages[:60]))
        if len(text.strip()) < 40 * max(1, len(reader.pages[:60])):      # a scan: look at the pages instead
            return f"{name} (scanned PDF, {len(reader.pages)} pages): " + _pdf_pages_seen(data, q)
    elif low.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        text = "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(
            " | ".join(c.text for c in row.cells) for t in d.tables for row in t.rows)
    elif low.endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        for ws in wb.worksheets[:8]:
            text += f"\n# sheet {ws.title}\n" + "\n".join(
                ",".join("" if v is None else str(v) for v in row) for row in ws.iter_rows(max_row=300, values_only=True))
    elif low.endswith(".pptx"):
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        text = "\n".join(f"[slide {i + 1}] " + " / ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
                         for i, s in enumerate(prs.slides))
    else:
        text = data[:400_000].decode("utf-8", errors="replace")
    text = text.strip()
    if not text:
        return f"{name}: I couldn't find any text in it."
    head = f"{name} ({len(text):,} characters)"
    return f"{head}:\n{text[:12000]}" + ("\n[the rest is longer; ask about a part]" if len(text) > 12000 else "")


def _pdf_pages_seen(data: bytes, q: str) -> str:
    try:
        import pypdfium2 as pdfium   # optional
    except ImportError:
        return "it has no text layer and I can't render pages here."
    import io
    pdf = pdfium.PdfDocument(data)
    imgs = []
    for i in range(min(len(pdf), 4)):
        out = io.BytesIO()
        pdf[i].render(scale=1.5).to_pil().save(out, "JPEG", quality=80)
        imgs.append(out.getvalue())
    return vision_describe(imgs, q)


if __name__ == "__main__":
    mcp.run()
