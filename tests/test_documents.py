"""Round-trip tests for the documents server (create each format, then read it back)."""
import json

import pytest
from mcp.server.mcpserver.exceptions import ToolError

from servers import documents as d

MD = """# Quarterly report
Sales were **strong** this quarter and *costs* fell.

## Highlights
- New office in Nairobi
- Hired 3 engineers

1. Plan budget
2. Review hires

| Item | Cost |
|------|-----:|
| Rent | 1200 |
| Food & drinks <bulk> | 300 |
"""


@pytest.fixture
def docs(tmp_path, monkeypatch):
    folder = tmp_path / "Documents"
    folder.mkdir()
    monkeypatch.setenv("IRIIS_ALLOWED_DIRS", json.dumps([str(folder)]))
    return folder


def test_parse_markdown():
    blocks = d.parse_markdown(MD)
    kinds = [b[0] for b in blocks]
    assert kinds == ["h", "p", "h", "ul", "ol", "table"]
    assert blocks[5][1] == [["Item", "Cost"], ["Rent", "1200"], ["Food & drinks <bulk>", "300"]]
    assert d._runs("a **b** *c*") == [("a ", False, False), ("b", True, False), (" ", False, False), ("c", False, True)]


def test_docx_round_trip(docs):
    out = d.create_document("Report", MD, title="Q3")
    assert out.startswith("Saved Report.docx in Documents")
    assert (docs / "Report.docx").is_file()
    text = d.read_document("Report.docx")
    for s in ("Q3", "Quarterly report", "Sales were strong", "New office in Nairobi", "Plan budget", "Rent\t1200"):
        assert s in text


def test_pdf_round_trip(docs):
    out = d.create_pdf("report.pdf", MD, title="Q3")
    assert "Saved report.pdf in Documents (1 page)" in out
    text = d.read_document("report.pdf")
    assert "[page 1]" in text and "Quarterly report" in text and "Food & drinks <bulk>" in text
    assert "Hired 3 engineers" in text


def test_xlsx_numbers_formulas_and_append(docs):
    out = d.create_spreadsheet("Budget", {"Costs": [["Item", "Cost"], ["Rent", 1200], ["Food", "300"],
                                                    ["Phone", "0712"], ["Total", "=SUM(B2:B3)"]],
                                          "Notes": [["Hello"]]})
    assert "Saved Budget.xlsx in Documents (2 sheets, 6 rows)" in out
    from openpyxl import load_workbook
    ws = load_workbook(docs / "Budget.xlsx")["Costs"]
    assert ws["B2"].value == 1200 and ws["B3"].value == 300 and ws["B4"].value == "0712"
    assert ws["B5"].value == "=SUM(B2:B3)" and ws["A1"].font.bold
    text = d.read_document("Budget.xlsx")
    assert "[sheet Costs]" in text and "Rent\t1200" in text and "=SUM(B2:B3)" in text and "[sheet Notes]" in text

    out = d.append_to_spreadsheet("Budget.xlsx", "Costs", [["Fuel", 50]])
    assert "Added 1 row to Costs" in out and "now 6 rows" in out
    out = d.append_to_spreadsheet("Budget.xlsx", "Extra", [["a", 1], ["b", 2]])
    assert "new sheet Extra" in out
    text = d.read_document("Budget.xlsx")
    assert "Fuel\t50" in text and "[sheet Extra]" in text
    with pytest.raises(ToolError):
        d.append_to_spreadsheet("Missing.xlsx", "Costs", [["x"]])


def test_pptx_round_trip(docs):
    out = d.create_presentation("Pitch", [{"title": "Iriis", "bullets": []},
                                          {"title": "Why", "bullets": ["Fast", "Private"], "notes": "Say hi"}])
    assert "Saved Pitch.pptx in Documents (2 slides)" in out
    text = d.read_document("Pitch.pptx")
    assert "[slide 1] Iriis" in text and "[slide 2] Why" in text and "Private" in text and "Notes: Say hi" in text


def test_csv_and_txt(docs):
    (docs / "a.csv").write_text("x,y\n1,2\n", encoding="utf-8")
    (docs / "b.md").write_text("hello there", encoding="utf-8")
    assert "x\ty\n1\t2" in d.read_document("a.csv")
    assert "hello there" in d.read_document("b.md")


def test_refuses_outside_allowed(docs, tmp_path):
    outside = tmp_path / "elsewhere" / "x.docx"
    with pytest.raises(ToolError, match="outside the allowed folders"):
        d.create_document(str(outside), "hi")
    with pytest.raises(ToolError, match="outside the allowed folders"):
        d.create_spreadsheet("../escape.xlsx", {"S": [[1]]})
    with pytest.raises(ToolError, match="outside the allowed folders"):
        d.read_document(str(tmp_path / "x.txt"))


def test_no_overwrite_without_flag(docs):
    d.create_document("Note", "first")
    with pytest.raises(ToolError, match="already exists"):
        d.create_document("Note.docx", "second")
    assert "first" in d.read_document("Note.docx")
    d.create_document("Note", "second", overwrite=True)
    assert "second" in d.read_document("Note.docx")
    d.create_pdf("Note", "x")
    with pytest.raises(ToolError, match="already exists"):
        d.create_pdf("Note", "y")


def test_extension_fixed_and_long_read_continues(docs):
    out = d.create_document("draft.txt", "word " * 3000)
    assert "draft.docx" in out
    first = d.read_document("draft.docx", max_chars=1000)
    assert "read_document again with start=1000" in first
    assert "more characters" not in d.read_document("draft.docx", max_chars=20000)
